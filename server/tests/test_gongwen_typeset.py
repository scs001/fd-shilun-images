"""公文国标排版与验收测试（tasks 3.1–3.6）：全离线。

测试口径：
- 版式断言从**产出的 docx** 反读（unzip + lxml + python-docx），不读 typeset 的常量；
- 字体断言遍历 document.xml 与两个 footer 的全部 run，任何一个漏写东亚字体槽都算失败；
- 国标清单逐项对照 gb_checklist()，不允许有要素既不落纸面也不进报告；
- PDF 用例在无 soffice 或无 pypdf 的机器上跳过（CI 不保证 LibreOffice）。
"""

from __future__ import annotations

import base64
import io
import json
import shutil
import zipfile
from collections import defaultdict

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from pydantic import ValidationError

from fd_shilun_api.gongwen import deliver, elements, fonts, typeset
from fd_shilun_api.gongwen.report import AcceptanceReport, CheckItem
from fd_shilun_api.locator import document_locator

TITLE = "关于进一步加强公文格式管理的通知"
BODY_PARAGRAPH = "一、统一版式标准。"
BODY_FRAGMENT = "统一版式标准"

FULL_FIELDS: dict = {
    "copy_number": "12",
    "secrecy": "秘密★1年",
    "urgency": "加急",
    "issuer_mark": "某某市人民政府文件",
    "doc_number": "国发〔2024〕5号",
    "signer": "张某某",
    "title": TITLE,
    "addressee": "各县（市、区）人民政府，市政府各部门",
    "body": "为进一步规范公文格式，现将有关事项通知如下：\n一、统一版式标准。\n二、加强监督检查。",
    "attachment_note": "附件：1.公文格式对照表",
    "attachments": ["公文格式对照表"],
    "issuer_signature": "某某市人民政府",
    "date": "2024年5月1日",
    "seal": True,
    "note": "此件公开发布",
    "cc": "市委办公室，市人大常委会办公室",
    "printing_org": "某某市人民政府办公室",
    "printing_date": "2024年5月2日",
}


def _full_doc(**overrides) -> elements.GongwenDocument:
    return elements.GongwenDocument(**{**FULL_FIELDS, **overrides})


def _minimal_doc(**overrides) -> elements.GongwenDocument:
    data = {
        "title": TITLE,
        "body": "只有正文，缺其余要素。",
        "issuer_signature": "某某市人民政府",
        "date": "2024年5月1日",
    }
    data.update(overrides)
    return elements.GongwenDocument(**data)


def _render(
    doc: elements.GongwenDocument | None = None,
    *,
    installed: set[str] | None = None,
    want_pdf: bool = False,
) -> deliver.DeliveryResult:
    """固定用「只装宋体」的环境渲染，保证不依赖本机字体。"""
    resolution = fonts.resolve_fonts(
        installed={"宋体"} if installed is None else installed
    )
    return deliver.render(
        doc if doc is not None else _full_doc(),
        want_pdf=want_pdf,
        fonts=resolution,
    )


def _zip_parts(docx_bytes: bytes) -> dict[str, str]:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        return {
            name: archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.endswith(".xml")
        }


def _part_text(xml: str) -> str:
    root = etree.fromstring(xml.encode("utf-8"))
    lines: list[str] = []
    for paragraph in root.iter(qn("w:p")):
        lines.append("".join(node.text or "" for node in paragraph.iter(qn("w:t"))))
    return "\n".join(lines)


def _all_text(docx_bytes: bytes) -> str:
    parts = _zip_parts(docx_bytes)
    chunks = [_part_text(parts["word/document.xml"])]
    for name in sorted(parts):
        if "footer" in name and name.startswith("word/"):
            chunks.append(_part_text(parts[name]))
    return "\n".join(chunks)


def _run_families(paragraph) -> set[str]:
    families: set[str] = set()
    for run in paragraph.runs:
        rpr = run._element.find(qn("w:rPr"))
        assert rpr is not None, f"run 缺少 rPr：{run.text!r}"
        rfonts = rpr.find(qn("w:rFonts"))
        assert rfonts is not None, f"run 缺少 rFonts：{run.text!r}"
        families.add(rfonts.get(qn("w:eastAsia")))
    return families


def _items_by_field(report: AcceptanceReport) -> dict[str, list[CheckItem]]:
    grouped: dict[str, list[CheckItem]] = defaultdict(list)
    for item in report.items:
        if item.locator.field:
            grouped[item.locator.field].append(item)
    return grouped


# ---------------------------------------------------------------------------
# 3.1 国标版式
# ---------------------------------------------------------------------------


def test_layout_parameters_measured_from_docx() -> None:
    result = _render()
    facts = typeset.layout_facts(io.BytesIO(result.docx))

    assert facts["page_size_mm"] == pytest.approx([210.0, 297.0], abs=0.1)
    assert facts["margins_mm"] == pytest.approx(
        {"top": 37.0, "left": 28.0, "bottom": 35.0, "right": 26.0}, abs=0.1
    )
    assert facts["text_area_mm"] == pytest.approx([156.0, 225.0], abs=0.1)
    assert facts["line_spacing_pt"] == pytest.approx(28.0, abs=0.01)
    assert facts["line_spacing_rule"] == "exact"
    assert facts["first_line_indent_twips"] == 640
    assert facts["first_line_indent_pt"] == pytest.approx(32.0, abs=0.01)
    # 22 行 × 28 字来自实测版心与字符网格，不是写死的常量
    assert facts["lines_per_page"] == 22
    assert facts["chars_per_line"] == 28

    # 再从原始 XML 复核一遍，避免只信任 python-docx 的换算
    parts = _zip_parts(result.docx)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    document_root = etree.fromstring(parts["word/document.xml"].encode("utf-8"))
    section = document_root.find(".//w:sectPr", ns)
    page_size = section.find("w:pgSz", ns)
    page_margins = section.find("w:pgMar", ns)
    assert page_size is not None and page_margins is not None
    assert int(page_size.get(qn("w:w"))) == pytest.approx(
        210 * 1440 / 25.4, abs=6
    )
    assert int(page_size.get(qn("w:h"))) == pytest.approx(
        297 * 1440 / 25.4, abs=6
    )
    for attr, expected_mm in (
        ("w:top", 37.0),
        ("w:left", 28.0),
        ("w:bottom", 35.0),
        ("w:right", 26.0),
    ):
        assert int(page_margins.get(qn(attr))) == pytest.approx(
            expected_mm * 1440 / 25.4, abs=6
        )

    styles = etree.fromstring(parts["word/styles.xml"].encode("utf-8"))
    normal = styles.find(".//w:style[@w:styleId='Normal']", ns)
    assert normal is not None
    spacing = normal.find("w:pPr/w:spacing", ns)
    indent = normal.find("w:pPr/w:ind", ns)
    assert int(spacing.get(qn("w:line"))) == 560
    assert spacing.get(qn("w:lineRule")) == "exact"
    assert int(indent.get(qn("w:firstLine"))) == 640


# ---------------------------------------------------------------------------
# 3.1 中文字体绑定（决策 3）
# ---------------------------------------------------------------------------


def test_every_run_binds_four_font_slots_and_east_asian_hint() -> None:
    result = _render(installed={"宋体"})
    parts = _zip_parts(result.docx)
    xml_parts = [parts["word/document.xml"]] + [
        parts[name]
        for name in sorted(parts)
        if name.startswith("word/footer") and name.endswith(".xml")
    ]
    assert len(xml_parts) >= 3  # 正文 + 奇偶页两个页脚

    run_count = 0
    for xml in xml_parts:
        root = etree.fromstring(xml.encode("utf-8"))
        for run in root.iter(qn("w:r")):
            run_count += 1
            rpr = run.find(qn("w:rPr"))
            assert rpr is not None, "存在没有 rPr 的 run"
            rfonts = rpr.find(qn("w:rFonts"))
            assert rfonts is not None, "存在没有 rFonts 的 run"
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                assert rfonts.get(qn(f"w:{attr}")), f"run 缺少 w:{attr} 字体槽"
            # w:hint 是 rFonts 的属性（CT_Fonts/@w:hint），必须为 eastAsia
            assert rfonts.get(qn("w:hint")) == "eastAsia"
    assert run_count >= 20


def test_title_and_body_use_first_available_family() -> None:
    installed = {"方正小标宋简体", "仿宋_GB2312", "黑体", "楷体_GB2312", "宋体"}
    result = _render(installed=installed)
    document = Document(io.BytesIO(result.docx))

    title_paragraph = next(p for p in document.paragraphs if p.text == TITLE)
    body_paragraph = next(p for p in document.paragraphs if p.text == BODY_PARAGRAPH)
    assert _run_families(title_paragraph) == {"方正小标宋简体"}
    assert _run_families(body_paragraph) == {"仿宋_GB2312"}

    font_items = {
        item.rule_id: item
        for item in result.report.items
        if item.rule_id.startswith("gov.font.")
    }
    for role in fonts.FONT_ROLES:
        item = font_items[f"gov.font.{role}"]
        assert item.status == "pass", item.evidence
    assert result.fonts.fallbacks == []


def test_only_songti_installed_falls_back_and_is_reported() -> None:
    result = _render(installed={"宋体"})
    document = Document(io.BytesIO(result.docx))

    title_paragraph = next(p for p in document.paragraphs if p.text == TITLE)
    body_paragraph = next(p for p in document.paragraphs if p.text == BODY_PARAGRAPH)
    assert _run_families(title_paragraph) == {"宋体"}
    assert _run_families(body_paragraph) == {"宋体"}

    font_items = {
        item.rule_id: item
        for item in result.report.items
        if item.rule_id.startswith("gov.font.")
    }
    title_item = font_items["gov.font.title"]
    body_item = font_items["gov.font.body"]
    assert title_item.status == "warn" and title_item.severity == "P1"
    assert "方正小标宋简体" in title_item.evidence and "宋体" in title_item.evidence
    assert body_item.status == "warn"
    assert "仿宋_GB2312" in body_item.evidence and "宋体" in body_item.evidence

    fallback_roles = {entry["role"] for entry in result.fonts.fallbacks}
    assert {"title", "body"} <= fallback_roles


# ---------------------------------------------------------------------------
# 3.2 版头/主体/版记/页码：逐项对照 18 项清单
# ---------------------------------------------------------------------------

# 18 项要素在完整样稿里的可查文本；None 表示该要素没有文字形态（由报告/页脚体现）
EXPECTED_TEXT: dict[str, str | None] = {
    "copy_number": "000012",
    "secrecy": "秘密★1年",
    "urgency": "加急",
    "issuer_mark": "某某市人民政府文件",
    "doc_number": "国发〔2024〕5号",
    "signer": "张某某",
    "title": TITLE,
    "addressee": "各县（市、区）人民政府",
    "body": BODY_FRAGMENT,
    "attachment_note": "附件：1.公文格式对照表",
    "issuer_signature": "某某市人民政府",
    "date": "2024年5月1日",
    "seal": None,
    "note": "（此件公开发布）",
    "attachments": "公文格式对照表",
    "cc": "抄送：市委办公室",
    "printing_org": "某某市人民政府办公室",
    "page_number": "— 1 —",
}


def test_gb_checklist_has_18_elements() -> None:
    checklist = elements.gb_checklist()
    assert len(checklist) == 18
    assert {entry["key"] for entry in checklist} == set(EXPECTED_TEXT)
    for entry in checklist:
        assert entry["label"]
        assert entry["gb_reference"].startswith("GB/T 9704-2012")
        assert isinstance(entry["required"], bool)
    required = {entry["key"] for entry in checklist if entry["required"]}
    assert required == {"title", "addressee", "body", "issuer_signature", "date"}


def test_full_document_accounts_for_every_checklist_item() -> None:
    result = _render()
    text = _all_text(result.docx)
    by_field = _items_by_field(result.report)

    for entry in elements.gb_checklist():
        key = entry["key"]
        expected = EXPECTED_TEXT[key]
        present = expected is not None and expected in text
        assert present or by_field[key], f"要素「{entry['label']}」既不在文档也不在报告"
        if not present:
            # 没有文字形态的要素必须在报告里有明确结论，不能静默
            assert any(item.status in ("pass", "warn", "fail") for item in by_field[key])

    # 完整样稿不应有 P0 失败
    assert result.report.deliverable is True
    assert result.report.counts["p0_fail"] == 0


def test_odd_and_even_footers_differ_and_carry_page_field() -> None:
    result = _render()
    facts = typeset.layout_facts(io.BytesIO(result.docx))["page_number"]
    assert facts["odd"] is True and facts["even"] is True
    assert facts["odd_alignment"] == "right"
    assert facts["even_alignment"] == "left"

    parts = _zip_parts(result.docx)
    footer_names = sorted(
        name for name in parts if name.startswith("word/footer") and name.endswith(".xml")
    )
    footer_xml = [parts[name] for name in footer_names]
    assert len(footer_xml) == 2
    assert footer_xml[0] != footer_xml[1]
    right_count = sum('<w:jc w:val="right"/>' in xml for xml in footer_xml)
    left_count = sum('<w:jc w:val="left"/>' in xml for xml in footer_xml)
    assert right_count == 1 and left_count == 1
    for xml in footer_xml:
        assert 'w:instr=" PAGE "' in xml
        assert "— " in xml and " —" in xml
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            assert f'w:{attr}="' in xml


def test_body_layout_order_head_main_record() -> None:
    result = _render()
    document = Document(io.BytesIO(result.docx))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # 版头 → 主体 → 版记 的段落顺序（不靠子串查找，避免「某某市人民政府」
    # 命中「某某市人民政府文件」这类前缀包含的歧义）
    assert paragraphs == [
        "000012",
        "秘密★1年",
        "加急",
        "某某市人民政府文件",
        "国发〔2024〕5号\t签发人：张某某",
        TITLE,
        "各县（市、区）人民政府，市政府各部门：",
        "为进一步规范公文格式，现将有关事项通知如下：",
        BODY_PARAGRAPH,
        "二、加强监督检查。",
        "附件：1.公文格式对照表",
        "某某市人民政府",
        "2024年5月1日",
        "（此件公开发布）",
        "抄送：市委办公室，市人大常委会办公室。",
        "某某市人民政府办公室\t2024年5月2日印发",
    ]

    # 成文日期右空四字、署名右空二字（GB 8.5）
    signature = next(p for p in document.paragraphs if p.text == "某某市人民政府")
    date = next(p for p in document.paragraphs if p.text == "2024年5月1日")
    assert signature.paragraph_format.right_indent.pt == pytest.approx(32.0)
    assert date.paragraph_format.right_indent.pt == pytest.approx(64.0)


# ---------------------------------------------------------------------------
# 3.3 发文字号与成文日期
# ---------------------------------------------------------------------------


def test_valid_doc_number_and_date_pass() -> None:
    result = _render()
    items = {item.rule_id: item for item in result.report.items}
    assert items["gov.doc_number_format"].status == "pass"
    assert items["gov.date_format"].status == "pass"
    assert result.report.deliverable is True


def test_invalid_doc_number_and_date_fail_with_correct_form() -> None:
    result = _render(
        _full_doc(doc_number="国发[2024]第05号", date="2024年05月01日")
    )
    items = {item.rule_id: item for item in result.report.items}

    doc_number = items["gov.doc_number_format"]
    assert doc_number.status == "fail" and doc_number.severity == "P0"
    assert doc_number.suggestion == "国发〔2024〕5号"
    assert "六角括号" in doc_number.evidence
    assert doc_number.locator.field == "doc_number"

    date = items["gov.date_format"]
    assert date.status == "fail" and date.severity == "P0"
    assert date.suggestion == "2024年5月1日"
    assert "虚位" in date.evidence

    assert result.report.deliverable is False
    assert result.report.counts["p0_fail"] == 2


# ---------------------------------------------------------------------------
# 3.3 必需要素与条件性要素
# ---------------------------------------------------------------------------


def test_partial_draft_marks_required_missing_and_optional_warn() -> None:
    result = _render(_minimal_doc())
    by_rule = {item.rule_id: item for item in result.report.items}

    for key in ("title", "addressee", "body", "issuer_signature", "date"):
        # 本用例样稿给了标题/正文/署名/日期，只缺主送机关
        if key == "addressee":
            item = by_rule["gov.element.addressee"]
            assert item.status == "fail" and item.severity == "P0"
            continue
        assert by_rule[f"gov.element.{key}"].status == "pass"

    for key in (
        "copy_number",
        "secrecy",
        "urgency",
        "signer",
        "attachment_note",
        "cc",
        "printing_org",
        "printing_date",
        "seal",
    ):
        item = by_rule[f"gov.element.{key}"]
        assert item.status == "warn", item.evidence
        assert item.severity == "P2", "条件性要素缺失不得判 P0"

    assert result.report.deliverable is False


def test_addressee_is_conditional_on_doctype() -> None:
    doc = _minimal_doc(issuer_signature="某某市人民政府", date="2024年5月1日")
    normal = {
        item.rule_id: item for item in elements.required_elements(doc)
    }["gov.element.addressee"]
    assert normal.status == "fail" and normal.severity == "P0"

    announcement = {
        item.rule_id: item
        for item in elements.required_elements(doc, doctype="公告")
    }["gov.element.addressee"]
    assert announcement.status == "warn" and announcement.severity == "P2"


def test_seal_and_attachment_consistency() -> None:
    doc = _full_doc(seal=False)
    items = {item.rule_id: item for item in elements.validate_elements(doc)}
    assert items["gov.element.seal_consistency"].status == "warn"
    assert items["gov.element.seal_consistency"].severity == "P1"

    mismatch = _full_doc(attachments=["公文格式对照表", "会议报名表"])
    items = {item.rule_id: item for item in elements.validate_elements(mismatch)}
    attachment = items["gov.element.attachment_mismatch"]
    assert attachment.status == "fail" and attachment.severity == "P1"
    assert "会议报名表" in attachment.evidence

    no_note = _full_doc(attachment_note=None)
    items = {item.rule_id: item for item in elements.validate_elements(no_note)}
    assert items["gov.element.attachment_mismatch"].status == "fail"


# ---------------------------------------------------------------------------
# 3.4 字体回退与严格模式
# ---------------------------------------------------------------------------


def test_font_resolution_fallback_and_strict_mode(monkeypatch) -> None:
    resolution = fonts.resolve_fonts(installed={"宋体"}, strict=False)
    assert resolution.family_for("title") == "宋体"
    assert resolution.family_for("body") == "宋体"
    assert "方正小标宋简体" in resolution.missing
    assert "仿宋_GB2312" in resolution.missing
    title_fallback = next(
        entry for entry in resolution.fallbacks if entry["role"] == "title"
    )
    assert title_fallback["requested"] == "方正小标宋简体"
    assert title_fallback["chosen"] == "宋体"
    assert "华文中宋" in title_fallback["missing"]

    with pytest.raises(fonts.MissingFontError) as excinfo:
        fonts.resolve_fonts(installed={"宋体"}, strict=True)
    assert "仿宋_GB2312" in str(excinfo.value)
    assert excinfo.value.missing

    monkeypatch.setenv("FD_SHILUN_FONT_STRICT", "1")
    with pytest.raises(fonts.MissingFontError):
        fonts.resolve_fonts(installed={"宋体"})
    monkeypatch.delenv("FD_SHILUN_FONT_STRICT")

    strict_ok = fonts.resolve_fonts(
        installed={"方正小标宋简体", "仿宋_GB2312", "黑体", "楷体_GB2312", "宋体"},
        strict=True,
    )
    assert strict_ok.fallbacks == []
    assert strict_ok.to_dict()["strict"] is True


def test_strict_mode_render_fails_instead_of_degrading(monkeypatch) -> None:
    monkeypatch.setattr(fonts, "available_fonts", lambda dirs=None: {"宋体"})
    with pytest.raises(fonts.MissingFontError):
        deliver.render(_full_doc(), strict=True)

    # 非严格模式仍产出 DOCX，但报告必须列出缺失与替代
    result = deliver.render(_full_doc(), strict=False)
    assert result.docx[:2] == b"PK"
    font_items = {
        item.rule_id: item
        for item in result.report.items
        if item.rule_id.startswith("gov.font.")
    }
    assert font_items["gov.font.title"].status == "warn"
    assert "方正小标宋简体" in font_items["gov.font.title"].evidence


# ---------------------------------------------------------------------------
# 3.5 验收报告
# ---------------------------------------------------------------------------


def test_acceptance_report_is_machine_readable() -> None:
    result = _render()
    payload = json.loads(result.report.to_json())

    assert payload["schema_version"] == "gongwen-report-v1"
    assert payload["doc_kind"] == "gongwen"
    assert payload["items"]
    assert result.report.counts["total"] == len(payload["items"])

    allowed_kinds = {"citation", "paragraph", "char_range", "document", "field"}
    for item in payload["items"]:
        assert item["rule_id"].strip()
        assert item["evidence"].strip()
        assert item["locator"]["kind"] in allowed_kinds
        assert item["status"] in {"pass", "fail", "warn"}

    assert result.report.pass_ and all(i.status == "pass" for i in result.report.pass_)
    assert result.report.failed == []
    assert result.report.deliverable is True

    with pytest.raises(ValidationError):
        CheckItem(
            rule_id="gov.blank_evidence",
            status="pass",
            locator=document_locator(),
            evidence="   ",
        )

    blocking = AcceptanceReport(
        items=[
            CheckItem(
                rule_id="gov.synthetic",
                status="fail",
                severity="P0",
                locator=document_locator(),
                evidence="构造的 P0 失败",
            )
        ]
    )
    assert blocking.deliverable is False
    assert blocking.counts["p0_fail"] == 1

    warning_only = AcceptanceReport(
        items=[
            CheckItem(
                rule_id="gov.synthetic",
                status="fail",
                severity="P1",
                locator=document_locator(),
                evidence="构造的 P1 失败",
            )
        ]
    )
    assert warning_only.deliverable is True


# ---------------------------------------------------------------------------
# API：base64 交付
# ---------------------------------------------------------------------------


def test_api_render_returns_report_and_base64() -> None:
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from fd_shilun_api.gongwen import api as gongwen_api

    app = FastAPI()
    app.include_router(gongwen_api.router)
    app.dependency_overrides[gongwen_api.require_user] = lambda: "user_a"
    client = TestClient(app)

    body = {**_full_doc().model_dump(), "want_pdf": False}
    response = client.post("/api/v1/gongwen/render", json=body)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["deliverable"] is True
    assert data["pdf_base64"] is None
    assert data["report"]["schema_version"] == "gongwen-report-v1"
    assert data["fonts"]["chosen"]["body"]

    docx_bytes = base64.b64decode(data["docx_base64"])
    assert docx_bytes[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        assert "word/document.xml" in archive.namelist()


# ---------------------------------------------------------------------------
# 3.6 PDF（可选）：中文可提取、字体已嵌入
# ---------------------------------------------------------------------------


def test_export_pdf_without_soffice_raises(monkeypatch) -> None:
    monkeypatch.setattr(deliver, "_find_soffice", lambda: None)
    with pytest.raises(deliver.PdfExportError) as excinfo:
        deliver.export_pdf(b"not-a-docx")
    assert "soffice" in str(excinfo.value)


@pytest.mark.skipif(
    shutil.which("soffice") is None, reason="未安装 LibreOffice（soffice），PDF 为可选交付"
)
def test_pdf_export_has_extractable_cjk_and_embedded_fonts() -> None:
    pytest.importorskip("pypdf", reason="pypdf 未安装，跳过 PDF 文本/字体检查")

    result = _render(want_pdf=True)
    assert result.pdf is not None
    assert result.pdf[:5] == b"%PDF-"

    text, embedded = deliver.pdf_text_and_fonts(result.pdf)
    compact = text.replace(" ", "").replace("\n", "")
    assert BODY_FRAGMENT in compact, f"PDF 中未提取到正文中文：{text[:200]!r}"
    # LibreOffice 会替换缺失字体，因此这里只要求「有嵌入字体」，不绑定商业字体名
    assert embedded, "PDF 未嵌入任何字体，中文可能显示为方框"

    assert result.report.deliverable is True
    assert any(item.rule_id == "gov.deliver.pdf" for item in result.report.items)
